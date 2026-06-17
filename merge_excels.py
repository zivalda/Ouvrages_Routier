"""
Fusion des fichiers selected_ouvrages.xlsx par département.

Structure attendue dans output_<route>/ :
    output_N4/
      51/
        T1_output_N4/selected_ouvrages.xlsx
        T2_output_N4/selected_ouvrages.xlsx
      54/
        T1_output_N4/selected_ouvrages.xlsx
      57/
        T1_output_N4/selected_ouvrages.xlsx

Produit un fichier fusionné par département :
    output_N4/
      51/51_N4_selected_ouvrages.xlsx
      54/54_N4_selected_ouvrages.xlsx
      57/57_N4_selected_ouvrages.xlsx

Usage :
    python merge_excels.py
    python merge_excels.py --route N4
    python merge_excels.py --route N4 --output-dir chemin/vers/output_N4
    python merge_excels.py --route N4 --dept 51        (un seul département)
"""

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ROOT_DIR = Path(__file__).parent

FILLS = {
    "remblai": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "deblai":  PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
    "rasant":  PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "mixte":   PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}


def troncon_sort_key(folder: Path) -> int:
    match = re.search(r"T(\d+)", folder.name)
    return int(match.group(1)) if match else 0


def merge_dept(dept: str, route: str, dept_output_dir: Path) -> None:
    """Fusionne tous les Excel d'un département et écrit le fichier final."""
    pattern = f"T*_output_{route}"
    troncon_folders = sorted(dept_output_dir.glob(pattern), key=troncon_sort_key)

    if not troncon_folders:
        print(f"  [AVERTISSEMENT] Aucun dossier '{pattern}' dans {dept_output_dir} — ignoré.")
        return

    print(f"\n  Département {dept} — {len(troncon_folders)} tronçon(s) :")

    frames = []
    missing = []

    for folder in troncon_folders:
        xlsx_path = folder / "selected_ouvrages.xlsx"
        if not xlsx_path.exists():
            print(f"    [MANQUANT] {folder.name}/selected_ouvrages.xlsx")
            missing.append(folder.name)
            continue

        df = pd.read_excel(xlsx_path, sheet_name="Ouvrages", engine="openpyxl")
        troncon_id = re.search(r"(T\d+)", folder.name).group(1)
        df.insert(0, "troncon", troncon_id)
        frames.append(df)
        print(f"    [OK] {folder.name} — {len(df)} ligne(s)")

    if not frames:
        print(f"    [ERREUR] Aucun Excel trouvé pour le département {dept}.")
        return

    merged = pd.concat(frames, ignore_index=True)
    if "id" in merged.columns:
        merged["id"] = range(1, len(merged) + 1)

    output_name = f"{dept}_{route}_selected_ouvrages.xlsx"
    output_path = dept_output_dir / output_name

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        merged.to_excel(writer, sheet_name="Ouvrages", index=False)
        ws = writer.sheets["Ouvrages"]

        # En-têtes : gras + centré
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Largeur automatique des colonnes
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        # Couleur par classification
        classification_col = None
        for cell in ws[1]:
            if str(cell.value).lower() == "classification":
                classification_col = cell.column
                break

        if classification_col is not None:
            for row in ws.iter_rows(min_row=2):
                classe = row[classification_col - 1].value
                fill = FILLS.get(str(classe).lower() if classe else "")
                if fill:
                    for cell in row:
                        cell.fill = fill

    print(f"    => {output_path.relative_to(ROOT_DIR)}  ({len(merged)} lignes au total)")
    if missing:
        print(f"    [ATTENTION] Tronçons sans Excel : {', '.join(missing)}")


def main():
    parser = argparse.ArgumentParser(
        description="Fusion des Excel selected_ouvrages par département"
    )
    parser.add_argument("--route",      type=str,  default=None, help="Code de la route (ex. N4)")
    parser.add_argument("--output-dir", type=Path, default=None, help="Dossier output_<route>/ (déduit si absent)")
    parser.add_argument("--dept",       type=str,  default=None, help="Traiter un seul département (ex. 51)")
    args = parser.parse_args()

    route = args.route or input("Code de la route (ex. N4) : ").strip()
    if not route:
        print("[ERREUR] Code route vide. Arrêt.")
        sys.exit(1)

    output_root: Path = args.output_dir or ROOT_DIR / f"output_{route}"
    if not output_root.exists():
        print(f"[ERREUR] Dossier introuvable : {output_root}")
        sys.exit(1)

    print(f"\nRoute   : {route}")
    print(f"Dossier : {output_root}")

    if args.dept:
        # Un seul département demandé
        dept_dir = output_root / args.dept
        if not dept_dir.exists():
            print(f"[ERREUR] Département {args.dept} introuvable dans {output_root}")
            sys.exit(1)
        merge_dept(args.dept, route, dept_dir)
    else:
        # Tous les départements présents dans output_root
        dept_dirs = sorted(
            [d for d in output_root.iterdir() if d.is_dir() and d.name.isdigit()],
            key=lambda d: int(d.name),
        )
        if not dept_dirs:
            print(f"[ERREUR] Aucun sous-dossier de département trouvé dans {output_root}")
            sys.exit(1)

        print(f"Départements détectés : {', '.join(d.name for d in dept_dirs)}\n")
        for dept_dir in dept_dirs:
            merge_dept(dept_dir.name, route, dept_dir)

    print("\n[OK] Fusion terminée.")


if __name__ == "__main__":
    main()
